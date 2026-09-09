import datetime
import secrets
import string
from io import BytesIO

import openpyxl

from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

from departments.models import Department
from courses.models import Course
from exams.models import Exam
from exam_questions.models import ExamQuestion
from results.models import Result

from .models import Student
from .forms import StudentForm


# ==========================================================
# PASSWORD GENERATOR
# ==========================================================

def generate_student_password(length=8):

    characters = string.ascii_letters + string.digits

    return "".join(
        secrets.choice(characters)
        for _ in range(length)
    )


# ==========================================================
# COMMON STUDENT CHECK
# ==========================================================

def get_logged_in_student(request):

    if not request.user.is_authenticated:
        return None

    try:
        return request.user.student
    except Student.DoesNotExist:
        return None


# ==========================================================
# COMMON EXAM TIME
# ==========================================================

def get_exam_times(exam):

    current_timezone = timezone.get_current_timezone()

    start_datetime = timezone.make_aware(
        datetime.datetime.combine(
            exam.exam_date,
            exam.start_time
        ),
        current_timezone
    )

    end_datetime = timezone.make_aware(
        datetime.datetime.combine(
            exam.exam_date,
            exam.end_time
        ),
        current_timezone
    )

    return start_datetime, end_datetime


# ==========================================================
# COMMON EXAM ELIGIBILITY
# ==========================================================

def student_is_eligible(student, exam):

    return (
        student.department_id == exam.department_id
        and student.course_id == exam.course_id
        and student.semester == exam.semester
    )


# ==========================================================
# CLEAR EXAM SESSION
# ==========================================================

def clear_exam_session(request, exam_id):

    request.session.pop(
        f"exam_ready_{exam_id}",
        None
    )

    request.session.pop(
        f"exam_ready_time_{exam_id}",
        None
    )

    request.session.pop(
        "active_exam_id",
        None
    )

    request.session.modified = True


# ==========================================================
# STUDENT LOGIN
# ==========================================================

def student_login(request):

    if request.user.is_authenticated:
        logout(request)

    if request.method == "POST":

        roll_number = request.POST.get(
            "roll_number",
            ""
        ).strip()

        password = request.POST.get(
            "password",
            ""
        )

        if not roll_number or not password:

            messages.error(
                request,
                "Please enter Roll Number and Password."
            )

            return redirect("student_login")

        user = authenticate(
            request,
            username=roll_number,
            password=password
        )

        if user is None:

            messages.error(
                request,
                "Invalid Roll Number or Password."
            )

            return redirect("student_login")

        if not hasattr(user, "student"):

            messages.error(
                request,
                "This account is not a student account."
            )

            return redirect("student_login")

        student = user.student

        if not student.is_active:

            messages.error(
                request,
                "Your student account is inactive."
            )

            return redirect("student_login")

        login(request, user)

        return redirect("student_dashboard")

    return render(
        request,
        "accounts/student_login.html"
    )


# ==========================================================
# BULK STUDENT UPLOAD
# ==========================================================

def student_bulk_upload(request):

    departments = Department.objects.all()
    courses = Course.objects.all()

    if request.method == "POST":

        department_id = request.POST.get("department")
        course_id = request.POST.get("course")
        semester_value = request.POST.get("semester")
        excel_file = request.FILES.get("excel_file")

        if not department_id:

            messages.error(
                request,
                "Please select a department."
            )

            return redirect("student_bulk_upload")

        if not course_id:

            messages.error(
                request,
                "Please select a course."
            )

            return redirect("student_bulk_upload")

        if not semester_value:

            messages.error(
                request,
                "Please select a semester."
            )

            return redirect("student_bulk_upload")

        if not excel_file:

            messages.error(
                request,
                "Please select an Excel file."
            )

            return redirect("student_bulk_upload")

        if not excel_file.name.lower().endswith(".xlsx"):

            messages.error(
                request,
                "Please upload an .xlsx Excel file."
            )

            return redirect("student_bulk_upload")

        department = get_object_or_404(
            Department,
            pk=department_id
        )

        course = get_object_or_404(
            Course,
            pk=course_id
        )

        try:

            semester = int(semester_value)

            if semester < 1 or semester > 10:
                raise ValueError

        except (TypeError, ValueError):

            messages.error(
                request,
                "Please select a valid semester."
            )

            return redirect("student_bulk_upload")

        try:

            workbook = openpyxl.load_workbook(
                excel_file,
                data_only=True
            )

            sheet = workbook.active

        except Exception:

            messages.error(
                request,
                "Unable to read the Excel file."
            )

            return redirect("student_bulk_upload")

        first_row = next(
            sheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True
            ),
            None
        )

        if not first_row:

            messages.error(
                request,
                "Excel file is empty."
            )

            return redirect("student_bulk_upload")

        headers = []

        for value in first_row:

            if value is None:
                headers.append("")
            else:
                headers.append(
                    str(value).strip().lower()
                )

        name_headers = {
            "name",
            "student name",
            "student_name"
        }

        roll_headers = {
            "roll number",
            "roll_number",
            "roll no",
            "rollno",
            "roll"
        }

        name_index = None
        roll_index = None

        for index, header in enumerate(headers):

            if header in name_headers:
                name_index = index

            if header in roll_headers:
                roll_index = index

        if name_index is None:

            messages.error(
                request,
                "Excel must contain a Name column."
            )

            return redirect("student_bulk_upload")

        if roll_index is None:

            messages.error(
                request,
                "Excel must contain a Roll Number column."
            )

            return redirect("student_bulk_upload")

        student_rows = []
        excel_roll_numbers = set()
        errors = []

        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True
            ),
            start=2
        ):

            name_value = (
                row[name_index]
                if name_index < len(row)
                else None
            )

            roll_value = (
                row[roll_index]
                if roll_index < len(row)
                else None
            )

            name = (
                str(name_value).strip()
                if name_value is not None
                else ""
            )

            roll_number = (
                str(roll_value).strip()
                if roll_value is not None
                else ""
            )

            if not name and not roll_number:
                continue

            if not name:

                errors.append(
                    f"Row {row_number}: Name is required."
                )

                continue

            if not roll_number:

                errors.append(
                    f"Row {row_number}: Roll Number is required."
                )

                continue

            roll_key = roll_number.upper()

            if roll_key in excel_roll_numbers:

                errors.append(
                    f"Row {row_number}: "
                    f"Duplicate Roll Number {roll_number}."
                )

                continue

            excel_roll_numbers.add(roll_key)

            if Student.objects.filter(
                roll_number__iexact=roll_number
            ).exists():

                errors.append(
                    f"Row {row_number}: "
                    f"Roll Number {roll_number} already exists."
                )

                continue

            if User.objects.filter(
                username__iexact=roll_number
            ).exists():

                errors.append(
                    f"Row {row_number}: "
                    f"Username {roll_number} already exists."
                )

                continue

            name_parts = name.split()

            first_name = name_parts[0]

            last_name = " ".join(
                name_parts[1:]
            )

            student_rows.append({
                "name": name,
                "roll_number": roll_number,
                "first_name": first_name,
                "last_name": last_name,
            })

        if errors:

            return render(
                request,
                "students/bulk_upload.html",
                {
                    "departments": departments,
                    "courses": courses,
                    "errors": errors,
                }
            )

        if not student_rows:

            messages.warning(
                request,
                "No students were found in the Excel file."
            )

            return redirect(
                "student_bulk_upload"
            )

        credentials = []

        try:

            with transaction.atomic():

                for data in student_rows:

                    roll_number = data["roll_number"]

                    password = generate_student_password()

                    user = User.objects.create_user(
                        username=roll_number,
                        password=password
                    )

                    Student.objects.create(
                        roll_number=roll_number,
                        first_name=data["first_name"],
                        last_name=data["last_name"],
                        email=None,
                        phone=None,
                        department=department,
                        course=course,
                        semester=semester,
                        is_active=True,
                        user=user
                    )

                    credentials.append({
                        "name": data["name"],
                        "roll_number": roll_number,
                        "password": password,
                        "department": str(department),
                        "course": str(course),
                        "semester": semester,
                    })

        except Exception as e:

            messages.error(
                request,
                f"Student upload failed: {e}"
            )

            return redirect(
                "student_bulk_upload"
            )

        request.session["student_credentials"] = credentials

        request.session["student_import_count"] = len(
            credentials
        )

        messages.success(
            request,
            f"{len(credentials)} students uploaded successfully."
        )

        return redirect(
            "student_upload_success"
        )

    return render(
        request,
        "students/bulk_upload.html",
        {
            "departments": departments,
            "courses": courses,
        }
    )


# ==========================================================
# STUDENT LIST
# ==========================================================

def student_list(request):

    students = Student.objects.select_related(
        "department",
        "course"
    ).all()

    return render(
        request,
        "students/student_list.html",
        {
            "students": students
        }
    )


# ==========================================================
# STUDENT CREATE
# ==========================================================

def student_create(request):

    form = StudentForm(
        request.POST or None
    )

    if form.is_valid():

        student = form.save()

        password = generate_student_password()

        user = User.objects.create_user(
            username=student.roll_number,
            password=password
        )

        student.user = user
        student.save()

        messages.success(
            request,
            f"Student created successfully. "
            f"Roll Number: {student.roll_number} | "
            f"Password: {password}"
        )

        return redirect(
            "student_list"
        )

    return render(
        request,
        "students/student_create.html",
        {
            "form": form
        }
    )


# ==========================================================
# STUDENT EDIT
# ==========================================================

def student_edit(request, pk):

    student = get_object_or_404(
        Student,
        pk=pk
    )

    form = StudentForm(
        request.POST or None,
        instance=student
    )

    if form.is_valid():

        form.save()

        messages.success(
            request,
            "Student updated successfully."
        )

        return redirect(
            "student_list"
        )

    return render(
        request,
        "students/student_edit.html",
        {
            "form": form,
            "student": student
        }
    )


# ==========================================================
# STUDENT DELETE
# ==========================================================

def student_delete(request, pk):

    student = get_object_or_404(
        Student,
        pk=pk
    )

    if request.method == "POST":

        user = student.user

        student.delete()

        if user:
            user.delete()

        messages.success(
            request,
            "Student deleted successfully."
        )

        return redirect(
            "student_list"
        )

    return render(
        request,
        "students/student_delete.html",
        {
            "student": student
        }
    )


# ==========================================================
# UPLOAD SUCCESS
# ==========================================================

def student_upload_success(request):

    credentials = request.session.get(
        "student_credentials",
        []
    )

    count = request.session.get(
        "student_import_count",
        0
    )

    return render(
        request,
        "students/upload_success.html",
        {
            "credentials": credentials,
            "count": count,
        }
    )


# ==========================================================
# DOWNLOAD STUDENT CREDENTIALS
# ==========================================================

def download_student_credentials(request):

    credentials = request.session.get(
        "student_credentials",
        []
    )

    if not credentials:

        messages.error(
            request,
            "No student credentials available."
        )

        return redirect(
            "student_list"
        )

    workbook = openpyxl.Workbook()

    sheet = workbook.active

    sheet.title = "Student Credentials"

    sheet.append([
        "Name",
        "Roll Number",
        "Password",
        "Department",
        "Course",
        "Semester",
    ])

    for student in credentials:

        sheet.append([
            student["name"],
            student["roll_number"],
            student["password"],
            student["department"],
            student["course"],
            student["semester"],
        ])

    widths = {
        "A": 30,
        "B": 20,
        "C": 20,
        "D": 25,
        "E": 25,
        "F": 12,
    }

    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="student_credentials.xlsx"'
    )

    return response


# ==========================================================
# STUDENT DASHBOARD
# ==========================================================

@login_required
def student_dashboard(request):

    student = get_logged_in_student(request)

    if not student:

        logout(request)

        messages.error(
            request,
            "Please login using a student account."
        )

        return redirect("student_login")

    exams = Exam.objects.filter(
        department=student.department,
        course=student.course,
        semester=student.semester,
        status="Active"
    )

    results = Result.objects.filter(
        student=student
    )

    today = timezone.localdate()

    upcoming_exams = exams.filter(
        exam_date__gte=today
    ).order_by(
        "exam_date",
        "start_time"
    )[:5]

    return render(
        request,
        "students/dashboard.html",
        {
            "student": student,
            "exam_count": exams.count(),
            "upcoming_exam_count": exams.filter(
                exam_date__gte=today
            ).count(),
            "result_count": results.count(),
            "upcoming_exams": upcoming_exams,
            "recent_results": results.order_by(
                "-id"
            )[:5],
        }
    )


# ==========================================================
# AVAILABLE EXAMS
# ==========================================================

@login_required
def student_exam_list(request):

    student = get_logged_in_student(request)

    if not student:

        logout(request)

        messages.error(
            request,
            "Please login using a student account."
        )

        return redirect("student_login")

    exams = Exam.objects.filter(
        department=student.department,
        course=student.course,
        semester=student.semester,
        status="Active"
    ).order_by(
        "exam_date",
        "start_time"
    )

    now = timezone.localtime()

    for exam in exams:

        start_datetime, end_datetime = get_exam_times(exam)

        existing_result = Result.objects.filter(
            student=student,
            exam=exam
        ).exists()

        exam.available = False
        exam.message = ""

        if existing_result:

            exam.message = "Already Submitted"

        elif now < start_datetime:

            exam.message = (
                f"Exam starts at "
                f"{exam.start_time.strftime('%I:%M %p')}"
            )

        elif now >= end_datetime:

            exam.message = "Exam Completed"

        else:

            exam.available = True
            exam.message = "Start Examination"

    return render(
        request,
        "students/exam_list.html",
        {
            "student": student,
            "exams": exams
        }
    )




# ==========================================================
# STUDENT RESULTS
# ==========================================================

@login_required
def student_results(request):

    student = get_logged_in_student(request)

    if not student:

        logout(request)

        messages.error(
            request,
            "Please login using a student account."
        )

        return redirect("student_login")

    results = (
        Result.objects
        .filter(
            student=student
        )
        .select_related(
            "exam"
        )
        .order_by(
            "-id"
        )
    )

    return render(
        request,
        "students/results.html",
        {
            "student": student,
            "results": results,
            "result": results.first(),
        }
    )


# ==========================================================
# STUDENT PROFILE
# ==========================================================

@login_required
def student_profile(request):

    student = get_logged_in_student(request)

    if not student:

        logout(request)

        messages.error(
            request,
            "Please login using a student account."
        )

        return redirect("student_login")

    return render(
        request,
        "students/profile.html",
        {
            "student": student
        }
    )


# ==========================================================
# STUDENT LOGOUT
# ==========================================================

def student_logout(request):

    logout(request)

    messages.success(
        request,
        "Logged out successfully."
    )

    return redirect("home")

@login_required
def student_exam_readiness(request, exam_id):

    student = get_logged_in_student(request)

    # ======================================================
    # STUDENT CHECK
    # ======================================================

    if not student:

        logout(request)

        messages.error(
            request,
            "Please login using a student account."
        )

        return redirect("student_login")

    # ======================================================
    # GET EXAM
    # ======================================================

    exam = get_object_or_404(
        Exam,
        id=exam_id,
        status="Active"
    )

    # ======================================================
    # ELIGIBILITY
    # ======================================================

    if (
        student.department_id != exam.department_id
        or student.course_id != exam.course_id
        or student.semester != exam.semester
    ):

        messages.error(
            request,
            "You are not eligible for this examination."
        )

        return redirect("student_exam_list")

    # ======================================================
    # ALREADY SUBMITTED
    # ======================================================

    if Result.objects.filter(
        student=student,
        exam=exam
    ).exists():

        messages.info(
            request,
            "You have already submitted this examination."
        )

        return redirect("student_results")

    # ======================================================
    # EXAM TIME
    # ======================================================

    start_datetime, end_datetime = get_exam_times(exam)

    now = timezone.localtime()

    if now < start_datetime:

        messages.warning(
            request,
            f"Exam starts at "
            f"{exam.start_time.strftime('%I:%M %p')}"
        )

        return redirect("student_exam_list")

    if now >= end_datetime:

        messages.error(
            request,
            "Exam time has expired."
        )

        return redirect("student_exam_list")

    # ======================================================
    # QUESTIONS CHECK
    # ======================================================

    questions = (
        ExamQuestion.objects
        .filter(exam=exam)
        .select_related("question")
        .order_by("question_order")
    )

    if not questions.exists():

        messages.error(
            request,
            "No questions are mapped to this examination."
        )

        return redirect("student_exam_list")

    # ======================================================
    # POST
    # CAMERA + MICROPHONE READY
    # ======================================================

    if request.method == "POST":

        camera_status = request.POST.get(
            "camera_status",
            ""
        ).strip().lower()

        microphone_status = request.POST.get(
            "microphone_status",
            ""
        ).strip().lower()

        # --------------------------------------------------
        # CAMERA
        # --------------------------------------------------

        if camera_status != "allowed":

            messages.error(
                request,
                "Camera permission is required."
            )

            return redirect(
                "student_exam_readiness",
                exam_id=exam.id
            )

        # --------------------------------------------------
        # MICROPHONE
        # --------------------------------------------------

        if microphone_status != "allowed":

            messages.error(
                request,
                "Microphone permission is required."
            )

            return redirect(
                "student_exam_readiness",
                exam_id=exam.id
            )

        # ==================================================
        # IMPORTANT
        # MARK EXAM AS READY
        # ==================================================

        request.session[f"exam_ready_{exam.id}"] = True

        request.session[
            f"exam_ready_time_{exam.id}"
        ] = timezone.now().isoformat()

        request.session["active_exam_id"] = exam.id

        request.session.modified = True

        # ==================================================
        # DIRECTLY OPEN ACTUAL EXAM
        # ==================================================

        return redirect(
            "student_start_exam",
            exam_id=exam.id
        )

    # ======================================================
    # GET
    # SHOW READINESS PAGE
    # ======================================================

    return render(
        request,
        "students/exam_readiness.html",
        {
            "exam": exam,
            "student": student,
            "start_datetime": start_datetime,
            "end_datetime": end_datetime,
            "questions": questions,
        }
    )

@login_required
def student_start_exam(request, exam_id):

    # ============================================================
    # GET LOGGED IN STUDENT
    # ============================================================

    student = get_logged_in_student(request)

    if not student:

        logout(request)

        messages.error(
            request,
            "Please login using a student account."
        )

        return redirect("student_login")


    # ============================================================
    # GET EXAM
    # ============================================================

    exam = get_object_or_404(
        Exam,
        id=exam_id,
        status="Active"
    )


    # ============================================================
    # CHECK STUDENT ELIGIBILITY
    # ============================================================

    if not student_is_eligible(student, exam):

        messages.error(
            request,
            "You are not eligible for this examination."
        )

        return redirect(
            "student_exam_list"
        )


    # ============================================================
    # CHECK ALREADY SUBMITTED
    # ============================================================

    existing_result = (
        Result.objects
        .filter(
            student=student,
            exam=exam
        )
        .first()
    )

    if existing_result:

        messages.info(
            request,
            "You have already submitted this examination."
        )

        return redirect(
            "student_results"
        )


    # ============================================================
    # EXAM TIME
    # ============================================================

    start_datetime, end_datetime = get_exam_times(exam)

    now = timezone.localtime()


    # ============================================================
    # EXAM NOT STARTED
    # ============================================================

    if now < start_datetime:

        messages.warning(
            request,
            (
                f"Exam starts at "
                f"{exam.start_time.strftime('%I:%M %p')}"
            )
        )

        return redirect(
            "student_exam_list"
        )


    # ============================================================
    # EXAM EXPIRED
    # ============================================================

    if now >= end_datetime:

        clear_exam_session(
            request,
            exam.id
        )

        messages.error(
            request,
            "The examination time has expired."
        )

        return redirect(
            "student_exam_list"
        )


    # ============================================================
    # GET EXAM QUESTIONS
    # ============================================================

    questions = (
        ExamQuestion.objects
        .filter(
            exam=exam
        )
        .select_related(
            "question"
        )
        .order_by(
            "question_order"
        )
    )


    # ============================================================
    # NO QUESTIONS
    # ============================================================

    if not questions.exists():

        clear_exam_session(
            request,
            exam.id
        )

        messages.error(
            request,
            "No questions are mapped to this examination."
        )

        return redirect(
            "student_exam_list"
        )


    # ============================================================
    # POST = SUBMIT EXAM
    # ============================================================

    if request.method == "POST":

        # --------------------------------------------------------
        # FINAL TIME CHECK
        # --------------------------------------------------------

        now = timezone.localtime()

        if now >= end_datetime:

            clear_exam_session(
                request,
                exam.id
            )

            messages.error(
                request,
                "The examination time has expired."
            )

            return redirect(
                "student_exam_list"
            )


        # ========================================================
        # CALCULATE MARKS
        # ========================================================

        total_marks = 0

        obtained_marks = 0


        for item in questions:

            question = item.question

            question_marks = question.marks or 0

            total_marks += question_marks

            answer = request.POST.get(
                f"question_{question.id}"
            )


            if answer == question.correct_answer:

                obtained_marks += question_marks


        # ========================================================
        # FALLBACK TOTAL MARKS
        # ========================================================

        if total_marks == 0:

            total_marks = exam.total_marks or 0


        # ========================================================
        # CALCULATE PERCENTAGE
        # ========================================================

        if total_marks > 0:

            percentage = round(
                (
                    obtained_marks /
                    total_marks
                ) * 100,
                2
            )

        else:

            percentage = 0


        # ========================================================
        # PASS / FAIL
        # ========================================================

        status = (
            "Pass"
            if obtained_marks >= exam.pass_marks
            else "Fail"
        )


        # ========================================================
        # SAVE RESULT
        # ========================================================

        try:

            with transaction.atomic():

                existing_result = (
                    Result.objects
                    .select_for_update()
                    .filter(
                        student=student,
                        exam=exam
                    )
                    .first()
                )


                # ------------------------------------------------
                # PREVENT DOUBLE SUBMISSION
                # ------------------------------------------------

                if existing_result:

                    clear_exam_session(
                        request,
                        exam.id
                    )

                    messages.info(
                        request,
                        "This examination has already been submitted."
                    )

                    return redirect(
                        "student_results"
                    )


                # ------------------------------------------------
                # CREATE RESULT
                # ------------------------------------------------

                Result.objects.create(

                    student=student,

                    exam=exam,

                    obtained_marks=obtained_marks,

                    total_marks=total_marks,

                    percentage=percentage,

                    status=status,
                )


        except Exception as e:

            messages.error(
                request,
                f"Unable to submit examination: {e}"
            )

            return redirect(
                "student_start_exam",
                exam_id=exam.id
            )


        # ========================================================
        # CLEAR EXAM SESSION
        # ========================================================

        clear_exam_session(
            request,
            exam.id
        )


        # ========================================================
        # SUCCESS
        # ========================================================

        messages.success(
            request,
            "Exam Submitted Successfully!"
        )


        return redirect(
            "student_results"
        )


    # ============================================================
    # GET = SHOW ACTUAL EXAM
    # ============================================================

    remaining_seconds = int(
        (
            end_datetime - now
        ).total_seconds()
    )


    # ============================================================
    # APPLY EXAM DURATION
    # ============================================================

    duration_seconds = (
        (exam.duration or 0) * 60
    )


    if duration_seconds > 0:

        remaining_seconds = min(
            remaining_seconds,
            duration_seconds
        )


    # ============================================================
    # STORE ACTIVE EXAM
    # ============================================================

    request.session["active_exam_id"] = exam.id

    request.session.modified = True


    # ============================================================
    # RENDER ACTUAL EXAM
    # ============================================================

    return render(
        request,
        "students/start_exam.html",
        {
            "exam": exam,

            "student": student,

            "questions": questions,

            "exam_questions": questions,

            "start_datetime": start_datetime,

            "end_datetime": end_datetime,

            "remaining_seconds": remaining_seconds,
        }
    )